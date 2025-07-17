#!/usr/bin/env python3
"""
Word文档导出服务

将生成的综述导出为Word文档格式
"""

import logging
import re
from typing import Dict, Any
from datetime import datetime
import tempfile
import os

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.shared import OxmlElement, qn
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

logger = logging.getLogger(__name__)


class WordExportService:
    """Word文档导出服务"""
    
    def __init__(self):
        if not DOCX_AVAILABLE:
            logger.warning("python-docx not available, Word export will be disabled")
    
    def export_review_to_word(self, review_data: Dict[str, Any]) -> str:
        """
        将综述导出为Word文档
        
        Args:
            review_data: 综述数据
            
        Returns:
            Word文档的文件路径
        """
        if not DOCX_AVAILABLE:
            raise Exception("python-docx not installed. Please install it with: pip install python-docx")
        
        try:
            # 创建新文档
            doc = Document()
            
            # 设置文档样式
            self._setup_document_styles(doc)
            
            # 添加标题
            title = review_data.get('title', '文献综述')
            title_paragraph = doc.add_heading(title, 0)
            title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # 添加生成信息
            info_paragraph = doc.add_paragraph()
            info_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            info_run = info_paragraph.add_run(f"生成时间：{datetime.now().strftime('%Y年%m月%d日')}")
            info_run.font.size = Pt(12)
            info_run.italic = True
            
            # 添加分隔线
            doc.add_paragraph("=" * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # 处理综述内容
            content = review_data.get('content', '')
            self._add_html_content_to_doc(doc, content)
            
            # 保存文档
            temp_dir = tempfile.gettempdir()
            filename = f"综述_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            filepath = os.path.join(temp_dir, filename)
            
            doc.save(filepath)
            
            logger.info(f"Word文档导出成功: {filepath}")
            return filepath
            
        except Exception as e:
            logger.error(f"Word文档导出失败: {e}")
            raise
    
    def _setup_document_styles(self, doc: Document):
        """设置文档样式"""
        try:
            # 设置正文样式
            styles = doc.styles
            
            # 正文样式
            if 'Normal' in styles:
                normal_style = styles['Normal']
                normal_font = normal_style.font
                normal_font.name = '宋体'
                normal_font.size = Pt(12)
                
            # 标题样式
            for i in range(1, 4):
                heading_name = f'Heading {i}'
                if heading_name in styles:
                    heading_style = styles[heading_name]
                    heading_font = heading_style.font
                    heading_font.name = '黑体'
                    heading_font.bold = True
                    if i == 1:
                        heading_font.size = Pt(16)
                    elif i == 2:
                        heading_font.size = Pt(14)
                    else:
                        heading_font.size = Pt(13)
                        
        except Exception as e:
            logger.warning(f"设置文档样式失败: {e}")
    
    def _add_html_content_to_doc(self, doc: Document, html_content: str):
        """将HTML内容添加到Word文档"""
        try:
            # 移除HTML标签并处理格式
            content = self._clean_html_content(html_content)
            
            # 按段落分割内容
            paragraphs = content.split('\n\n')
            
            for para_text in paragraphs:
                para_text = para_text.strip()
                if not para_text:
                    continue
                
                # 检查是否是标题
                if para_text.startswith('# '):
                    # 一级标题
                    title_text = para_text[2:].strip()
                    doc.add_heading(title_text, 1)
                elif para_text.startswith('## '):
                    # 二级标题
                    title_text = para_text[3:].strip()
                    doc.add_heading(title_text, 2)
                elif para_text.startswith('### '):
                    # 三级标题
                    title_text = para_text[4:].strip()
                    doc.add_heading(title_text, 3)
                else:
                    # 普通段落
                    paragraph = doc.add_paragraph()
                    
                    # 处理引用格式 [1], [2] 等
                    parts = re.split(r'(\[\d+\])', para_text)
                    
                    for part in parts:
                        if re.match(r'\[\d+\]', part):
                            # 这是引用
                            run = paragraph.add_run(part)
                            run.font.superscript = True
                            run.font.color.rgb = None  # 使用默认颜色
                        else:
                            # 普通文本
                            if part.strip():
                                paragraph.add_run(part)
                    
                    # 设置段落格式
                    paragraph.paragraph_format.first_line_indent = Inches(0.5)
                    paragraph.paragraph_format.line_spacing = 1.5
                    
        except Exception as e:
            logger.error(f"添加HTML内容到Word文档失败: {e}")
            # 如果处理失败，添加原始内容
            doc.add_paragraph(html_content)
    
    def _clean_html_content(self, html_content: str) -> str:
        """清理HTML内容，转换为纯文本"""
        try:
            # 替换HTML标题标签
            content = re.sub(r'<h1[^>]*>(.*?)</h1>', r'# \1', html_content, flags=re.IGNORECASE | re.DOTALL)
            content = re.sub(r'<h2[^>]*>(.*?)</h2>', r'## \1', content, flags=re.IGNORECASE | re.DOTALL)
            content = re.sub(r'<h3[^>]*>(.*?)</h3>', r'### \1', content, flags=re.IGNORECASE | re.DOTALL)
            content = re.sub(r'<h[4-6][^>]*>(.*?)</h[4-6]>', r'### \1', content, flags=re.IGNORECASE | re.DOTALL)
            
            # 替换段落标签
            content = re.sub(r'<p[^>]*>(.*?)</p>', r'\1\n\n', content, flags=re.IGNORECASE | re.DOTALL)
            
            # 替换列表标签
            content = re.sub(r'<ol[^>]*>(.*?)</ol>', r'\1', content, flags=re.IGNORECASE | re.DOTALL)
            content = re.sub(r'<ul[^>]*>(.*?)</ul>', r'\1', content, flags=re.IGNORECASE | re.DOTALL)
            content = re.sub(r'<li[^>]*>(.*?)</li>', r'• \1\n', content, flags=re.IGNORECASE | re.DOTALL)
            
            # 替换强调标签
            content = re.sub(r'<strong[^>]*>(.*?)</strong>', r'\1', content, flags=re.IGNORECASE | re.DOTALL)
            content = re.sub(r'<b[^>]*>(.*?)</b>', r'\1', content, flags=re.IGNORECASE | re.DOTALL)
            content = re.sub(r'<em[^>]*>(.*?)</em>', r'\1', content, flags=re.IGNORECASE | re.DOTALL)
            content = re.sub(r'<i[^>]*>(.*?)</i>', r'\1', content, flags=re.IGNORECASE | re.DOTALL)
            
            # 移除其他HTML标签
            content = re.sub(r'<[^>]+>', '', content)
            
            # 清理多余的空行
            content = re.sub(r'\n\s*\n\s*\n', '\n\n', content)
            
            # 解码HTML实体
            content = content.replace('&nbsp;', ' ')
            content = content.replace('&lt;', '<')
            content = content.replace('&gt;', '>')
            content = content.replace('&amp;', '&')
            content = content.replace('&quot;', '"')
            
            return content.strip()
            
        except Exception as e:
            logger.error(f"清理HTML内容失败: {e}")
            return html_content


# 全局服务实例
word_export_service = WordExportService()
