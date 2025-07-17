"""
导出API路由 - 完全复制原版本实现
"""

from flask import Blueprint, request, jsonify, send_file, session
from io import BytesIO
import logging
from datetime import datetime
import re

from utils.session_cache import get_papers_cache
from original_export_functions import generate_excel_in_memory, generate_word_in_memory

logger = logging.getLogger(__name__)

# Create blueprint
export_bp = Blueprint('export', __name__, url_prefix='/api/export')


@export_bp.route('/excel', methods=['POST'])
def export_excel():
    """导出Excel文件API端点 - 完全复制原版本实现"""
    try:
        data = request.get_json()
        query = data.get('query', '').strip()

        # 获取会话ID
        session_id = request.headers.get('sid') or session.get('session_id')
        if not session_id:
            return jsonify({'error': '无效的会话ID'}), 400

        # 从缓存中获取文献数据
        cache = get_papers_cache()
        cached_data = cache.get_papers(session_id)
        if not cached_data:
            return jsonify({'error': '文献数据不存在，请重新检索'}), 404

        # 根据data_type决定使用哪种数据
        data_type = data.get('data_type', 'filtered')
        if data_type == 'initial' and 'original_papers' in cached_data:
            papers_data = {
                'papers': cached_data['original_papers'],
                'query': cached_data['query']
            }
        else:
            papers_data = {
                'papers': cached_data['papers'],
                'query': cached_data['query']
            }

        # 生成Excel文件
        file_content, filename = generate_excel_in_memory(papers_data)

        if file_content:
            # 创建BytesIO对象
            excel_buffer = BytesIO(file_content)
            excel_buffer.seek(0)
            # 手动设置headers确保兼容性
            from flask import make_response
            response = make_response(file_content)
            response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
            response.headers['Content-Length'] = len(file_content)
            return response
        else:
            return jsonify({'error': '生成Excel文件失败'}), 500

    except Exception as e:
        logger.error(f"导出Excel文件时发生错误: {str(e)}")
        return jsonify({'error': str(e)}), 500


@export_bp.route('/word', methods=['POST'])
def export_word():
    """导出Word文件API端点 - 完全复制原版本实现"""
    try:
        data = request.get_json()
        query = data.get('query', '').strip()

        # 获取会话ID
        session_id = request.headers.get('sid') or session.get('session_id')
        if not session_id:
            return jsonify({'error': '无效的会话ID'}), 400

        # 从缓存中获取文献数据
        cache = get_papers_cache()
        cached_data = cache.get_papers(session_id)
        if not cached_data:
            return jsonify({'error': '文献数据不存在，请重新检索'}), 404

        # 根据data_type决定使用哪种数据
        data_type = data.get('data_type', 'filtered')
        if data_type == 'initial' and 'original_papers' in cached_data:
            papers_data = {
                'papers': cached_data['original_papers'],
                'query': cached_data['query']
            }
        else:
            papers_data = {
                'papers': cached_data['papers'],
                'query': cached_data['query']
            }

        # 生成Word文件
        file_content, filename = generate_word_in_memory(papers_data)

        if file_content:
            # 创建BytesIO对象
            word_buffer = BytesIO(file_content)
            word_buffer.seek(0)
            # 手动设置headers确保兼容性
            from flask import make_response
            response = make_response(file_content)
            response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
            response.headers['Content-Length'] = len(file_content)
            return response
        else:
            return jsonify({'error': '生成Word文件失败'}), 500

    except Exception as e:
        logger.error(f"导出Word文件时发生错误: {str(e)}")
        return jsonify({'error': str(e)}), 500


@export_bp.route('/review_word', methods=['POST'])
def export_review_word():
    """导出综述Word文档"""
    try:
        data = request.get_json()
        if not data:
            return jsonify({'error': '缺少请求数据'}), 400

        title = data.get('title', '文献综述')
        content = data.get('content', '')
        references = data.get('references', [])

        if not content:
            return jsonify({'error': '综述内容为空'}), 400

        # 生成Word文档
        word_content = generate_review_word_document(title, content, references)

        if word_content:
            # 创建响应
            from flask import make_response
            response = make_response(word_content)
            response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

            # 生成文件名
            safe_title = re.sub(r'[^\w\s-]', '', title).strip()[:50]
            filename = f"{safe_title}_综述_{datetime.now().strftime('%Y%m%d')}.docx"
            response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
            response.headers['Content-Length'] = len(word_content)

            logger.info(f"成功导出综述Word文档: {filename}")
            return response
        else:
            return jsonify({'error': '生成Word文档失败'}), 500

    except Exception as e:
        logger.error(f"导出综述Word文档时发生错误: {str(e)}")
        return jsonify({'error': str(e)}), 500


def generate_review_word_document(title: str, content: str, references: list) -> bytes:
    """
    生成综述Word文档

    Args:
        title: 综述标题
        content: 综述内容（HTML格式）
        references: 参考文献列表

    Returns:
        Word文档的字节内容
    """
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.style import WD_STYLE_TYPE
        import re
        from io import BytesIO

        # 创建文档
        doc = Document()

        # 设置页面边距
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1.25)
            section.right_margin = Inches(1.25)

        # 添加标题
        title_paragraph = doc.add_heading(title, 0)
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 处理HTML内容
        html_content = content

        # 移除CSS样式
        html_content = re.sub(r'<style>.*?</style>', '', html_content, flags=re.DOTALL)

        # 处理HTML标签
        sections = re.split(r'<h[12]>', html_content)

        for section in sections:
            if not section.strip():
                continue

            # 提取标题
            title_match = re.search(r'^(.*?)</h[12]>', section)
            if title_match:
                section_title = title_match.group(1).strip()
                section_content = section[title_match.end():]

                # 添加章节标题
                if section_title:
                    heading = doc.add_heading(section_title, level=1)
                    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                section_content = section

            # 处理段落内容
            paragraphs = re.split(r'</?p>', section_content)
            for para in paragraphs:
                para = para.strip()
                if not para:
                    continue

                # 清理HTML标签
                para = re.sub(r'<[^>]+>', '', para)
                para = para.strip()

                if para:
                    p = doc.add_paragraph(para)
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        # 添加参考文献
        if references:
            doc.add_page_break()
            ref_heading = doc.add_heading('参考文献', level=1)
            ref_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

            for i, ref in enumerate(references, 1):
                ref_para = doc.add_paragraph(f"{i}. {ref}")
                ref_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        # 保存到内存
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        return buffer.getvalue()

    except ImportError:
        logger.error("python-docx库未安装，无法生成Word文档")
        return None
    except Exception as e:
        logger.error(f"生成综述Word文档失败: {e}")
        return None