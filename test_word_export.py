#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
测试Word导出功能
"""

import sys
import os
from io import BytesIO
from datetime import datetime
import re

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    print("✅ python-docx 库导入成功")
except ImportError as e:
    print(f"❌ python-docx 库导入失败: {e}")
    sys.exit(1)

def test_word_generation():
    """测试Word文档生成"""
    try:
        # 模拟文献数据
        papers_data = {
            'query': '测试查询',
            'papers': [
                {
                    'title': '测试文献标题1',
                    'authors': ['作者1', '作者2'],
                    'journal_info': {
                        'title': '测试期刊',
                        'impact_factor': 5.234,
                        'jcr_quartile': 1,
                        'cas_quartile': 2
                    },
                    'pub_year': '2023',
                    'relevance': 95.5,
                    'pmid': '12345678',
                    'abstract': '这是一个测试摘要，用于验证Word导出功能是否正常工作。'
                },
                {
                    'title': '测试文献标题2',
                    'authors': ['作者3', '作者4'],
                    'journal_info': {
                        'title': '另一个测试期刊',
                        'impact_factor': 3.456,
                        'jcr_quartile': 2,
                        'cas_quartile': 3
                    },
                    'pub_year': '2024',
                    'relevance': 88.2,
                    'pmid': '87654321',
                    'abstract': '这是第二个测试摘要。'
                }
            ]
        }
        
        papers = papers_data.get('papers', [])
        query = papers_data.get('query', '未知查询')

        print(f"📝 开始生成Word文档，包含 {len(papers)} 篇文献...")

        # 创建Word文档
        doc = Document()

        # 添加标题
        title = doc.add_heading(f'文献检索报告 - {query}', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 添加基本信息
        doc.add_heading('检索信息', level=1)
        info_para = doc.add_paragraph()
        info_para.add_run('检索主题: ').bold = True
        info_para.add_run(query)
        info_para.add_run('\n检索时间: ').bold = True
        info_para.add_run(datetime.now().strftime("%Y年%m月%d日 %H:%M:%S"))
        info_para.add_run('\n文献数量: ').bold = True
        info_para.add_run(str(len(papers)))

        # 添加文献列表
        doc.add_heading('文献列表', level=1)

        for i, paper in enumerate(papers, 1):
            # 文献标题
            title_para = doc.add_paragraph()
            title_para.add_run(f'{i}. ').bold = True
            title_para.add_run(paper.get('title', '无标题')).bold = True

            # 文献信息
            info_para = doc.add_paragraph()
            authors = ', '.join(paper.get('authors', []))
            if authors:
                info_para.add_run('作者: ').bold = True
                info_para.add_run(authors + '\n')

            journal = paper.get('journal_info', {}).get('title', '')
            if journal:
                info_para.add_run('期刊: ').bold = True
                info_para.add_run(journal + '\n')

            # 影响因子
            impact_factor = paper.get('journal_info', {}).get('impact_factor', '')
            if impact_factor:
                info_para.add_run('影响因子: ').bold = True
                info_para.add_run(str(impact_factor) + '\n')

            # JCR分区
            jcr_quartile = paper.get('journal_info', {}).get('jcr_quartile', '')
            if jcr_quartile:
                info_para.add_run('JCR分区: ').bold = True
                info_para.add_run(f'Q{jcr_quartile}\n')

            # 中科院分区
            cas_quartile = paper.get('journal_info', {}).get('cas_quartile', '')
            if cas_quartile:
                info_para.add_run('中科院分区: ').bold = True
                info_para.add_run(f'{cas_quartile}区\n')

            pub_year = paper.get('pub_year', '')
            if pub_year:
                info_para.add_run('年份: ').bold = True
                info_para.add_run(str(pub_year) + '\n')

            relevance = paper.get('relevance', 0)
            info_para.add_run('相关度: ').bold = True
            info_para.add_run(f'{relevance:.1f}%\n')

            pmid = paper.get('pmid', '')
            if pmid:
                info_para.add_run('PMID: ').bold = True
                info_para.add_run(pmid + '\n')

            # 摘要
            abstract = paper.get('abstract', '')
            if abstract:
                abstract_para = doc.add_paragraph()
                abstract_para.add_run('摘要: ').bold = True
                abstract_para.add_run(abstract)

            # 添加分隔线
            if i < len(papers):
                doc.add_paragraph('─' * 50)

        # 保存到内存
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        file_content = output.getvalue()

        # 生成文件名
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        clean_query = re.sub(r'[^\w\s-]', '', query)[:20]
        filename = f"test_report_{clean_query}_{timestamp}.docx"

        print(f"✅ Word文档生成成功！")
        print(f"📄 文件大小: {len(file_content)} 字节")
        print(f"📁 文件名: {filename}")

        # 保存到文件进行验证
        with open(filename, 'wb') as f:
            f.write(file_content)
        
        print(f"💾 测试文件已保存到: {os.path.abspath(filename)}")
        
        return file_content, filename

    except Exception as e:
        print(f"❌ Word文档生成失败: {e}")
        import traceback
        traceback.print_exc()
        return None, None

if __name__ == "__main__":
    print("🧪 开始测试Word导出功能...")
    test_word_generation()
    print("🏁 测试完成")
